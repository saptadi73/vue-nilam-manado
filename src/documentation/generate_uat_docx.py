from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).with_name("Dokumentasi_UAT_NILAM_Manado.docx")


def cell_text(cell, text, bold=False, color=None, size=8):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None, size=8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, value in enumerate(headers):
        cell_text(t.rows[0].cells[i], value, True, "FFFFFF", size)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "176B52")
        t.rows[0].cells[i]._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, size=size)
            if widths:
                cells[i].width = Cm(widths[i])
    return t


def heading(doc, value, level=1):
    p = doc.add_heading(value, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, value):
    doc.add_paragraph(value, style="List Bullet")


raw_cases = [
    ("Publik", "Akses beranda dan Tentang", "Belum login", "Buka /home dan /tentang; gunakan navigasi.", "Halaman tampil tanpa autentikasi dan tanpa error.", "-"),
    ("Autentikasi", "Proteksi route Real API", "Belum login", "Buka /real/dashboard secara langsung.", "Dialihkan ke login dan tujuan redirect dipertahankan.", "Router guard"),
    ("Autentikasi", "Login valid", "Akun aktif tersedia", "Isi email/password valid lalu Login.", "Token tersimpan dan pengguna masuk ke dashboard/tujuan.", "POST /auth/login"),
    ("Autentikasi", "Login invalid", "Berada di login", "Masukkan kredensial salah.", "Pesan error tampil dan token tidak tersimpan.", "POST /auth/login"),
    ("Autentikasi", "Sesi kedaluwarsa dan logout", "Sudah login", "Gunakan token expired atau logout lalu akses route terproteksi.", "Token dibersihkan dan kembali ke login.", "401 handling"),
    ("Dashboard", "Ringkasan, grafik, dan peringkat", "Login; data transaksi tersedia", "Buka /real/dashboard dan tunggu seluruh widget.", "Agregat penjualan, biaya, produksi, kabupaten, dan net profit sesuai data.", "GET /dashboard/*"),
    ("Dashboard", "Filter tanggal dan petani", "Data lintas tanggal/petani tersedia", "Atur rentang tanggal dan petani; terapkan lalu reset.", "Semua widget mengikuti filter; reset mengembalikan seluruh data.", "GET /dashboard/*"),
    ("Dashboard", "Validasi rentang tanggal", "Login", "Pilih tanggal akhir sebelum tanggal mulai.", "Input ditolak atau pesan validasi yang jelas tampil.", "Dashboard query"),
    ("Dashboard", "Cari dan paginasi peringkat", "Data melebihi satu halaman", "Cari nama/NIK/wilayah dan ubah halaman.", "Hasil, urutan, jumlah, dan halaman konsisten.", "GET /dashboard/farmer-net-profit"),
    ("Profil", "Tampil dan ubah profil", "Login", "Buka /real/profile, ubah field yang diizinkan, simpan dan refresh.", "Data profil benar dan perubahan persisten.", "User/profile API"),
    ("Wilayah", "Hierarki dan pencarian wilayah", "Login", "Pilih/cari provinsi hingga desa; ubah induk.", "Anak sesuai induk dan pilihan turunan direset saat induk berubah.", "GET /wilayah/*"),
    ("Petani", "List, cari, dan paginasi", "Data tersedia", "Buka /real/petani; cari nama/wilayah; ubah halaman.", "Hasil dan paginasi konsisten.", "GET /farmers"),
    ("Petani", "Create petani valid dengan foto", "NIK unik; wilayah tersedia", "Isi field wajib, NIK 16 digit, kontak/wilayah, pilih foto valid, simpan.", "Preview tampil dan data/foto tersimpan.", "POST /farmers/with-foto"),
    ("Petani", "Validasi data invalid", "Login", "Coba NIK bukan 16 digit/berhuruf/duplikat, HP invalid, wilayah tidak konsisten.", "Invalid ditolak dengan pesan field; tidak ada record baru.", "POST/PUT /farmers"),
    ("Petani", "Detail, edit, foto, hapus", "Petani tersedia", "Buka detail; edit; ganti/hapus foto; uji batal/setuju hapus.", "Perubahan persisten; konfirmasi dan dependensi hapus ditangani.", "/farmers/{id}"),
    ("Petani", "Ringkasan petani", "Petani memiliki transaksi", "Buka summary, lahan, expense, dan produksi terkait.", "Semua total dan record hanya milik petani terpilih.", "GET /dashboard/farmers/{id}/summary"),
    ("Lahan", "List, cari, detail, paginasi", "Data lahan tersedia", "Cari kode/pemilik/wilayah; buka detail; ubah halaman.", "Hasil benar; detail menampilkan peta, foto, dan wilayah.", "GET /lands"),
    ("Lahan", "Create lahan valid", "Petani/wilayah tersedia", "Isi kode unik, luas >0, kepemilikan, wilayah, >=3 titik valid dan foto; simpan.", "Lahan, polygon, foto, dan relasi tersimpan.", "POST /lands"),
    ("Lahan", "Validasi lahan invalid", "Login", "Coba kode duplikat, luas <=0, wilayah parsial, koordinat tidak berpasangan/<3/di luar range, foto >5MB/non-gambar.", "Invalid ditolak dan tidak tersimpan.", "POST/PUT /lands"),
    ("Lahan", "Edit, foto, dan hapus", "Lahan tersedia", "Edit data/polygon; ganti/hapus foto; uji konfirmasi hapus.", "Perubahan persisten dan hapus aman.", "/lands/{id}"),
    ("Mitra", "List, cari, filter dan CRUD", "Wilayah tersedia", "Cari/filter; buat mitra valid; edit; hapus.", "CRUD, paginasi, wilayah, kontak, dan audit benar.", "/partners"),
    ("Mitra", "Validasi mitra invalid", "Login", "Coba nama/alamat kosong, email/HP/web invalid, wilayah tidak lengkap/konsisten.", "Input invalid ditolak.", "POST/PUT /partners"),
    ("Produk Biaya", "List, cari dan CRUD", "Login", "Buat nama, harga >=0, satuan; cari; edit; hapus.", "CRUD dan paginasi berhasil.", "/financing-products"),
    ("Produk Biaya", "Validasi invalid", "Login", "Kosongkan nama/satuan atau isi harga negatif/nonangka.", "Penyimpanan ditolak.", "POST/PUT /financing-products"),
    ("Produk Penjualan", "List, filter jenis dan CRUD", "Login", "Cari/filter barang-jasa; buat; edit; hapus.", "CRUD, filter, dan paginasi berhasil.", "/sales-products"),
    ("Produk Penjualan", "Validasi invalid", "Login", "Kosongkan nama/satuan, jenis di luar barang/jasa, atau harga negatif.", "Input invalid ditolak.", "POST/PUT /sales-products"),
    ("Produksi Tanam", "List, cari, filter dan paginasi", "Data tersedia", "Cari kode/petani/lahan; filter status/petani/lahan; ubah halaman.", "Hasil dan total sesuai kombinasi filter.", "GET /planting-productions"),
    ("Produksi Tanam", "Create valid", "Petani dan lahan miliknya tersedia", "Isi kode unik, tanggal, status, petani/lahan, luas dan angka rencana valid.", "Data dibuat dengan relasi dan audit benar.", "POST /planting-productions"),
    ("Produksi Tanam", "Validasi dan aturan selesai", "Produksi tersedia", "Coba kode invalid, lahan petani lain, angka negatif, tanggal akhir < mulai, atau status selesai tanpa hasil/tanggal aktual.", "Invalid ditolak; record valid dapat diselesaikan.", "POST/PUT /planting-productions"),
    ("Produksi Tanam", "Detail, edit, hapus dan catatan", "Produksi tersedia", "Buka/edit/hapus; tambah/edit/hapus catatan; coba catatan kosong/>2000 karakter.", "CRUD produksi dan catatan bekerja serta tervalidasi.", "/planting-productions/{id}"),
    ("Produksi Minyak", "List, cari, filter dan paginasi", "Data tersedia", "Cari dan filter status/petani/lahan; ubah halaman.", "Hasil dan total sesuai filter.", "GET /oil-productions"),
    ("Produksi Minyak", "Create valid", "Petani/lahan tersedia", "Isi kode, tanggal, status, relasi, berat/rencana dan field kilang; simpan.", "Data dibuat dan muncul pada list/detail.", "POST /oil-productions"),
    ("Produksi Minyak", "Validasi, selesai, dan rendemen", "Produksi tersedia", "Coba nilai negatif/relasi salah/tanggal invalid; selesaikan tanpa field aktual; lalu lengkapi valid.", "Invalid ditolak; valid tersimpan; rendemen = hasil minyak/berat kering.", "POST/PUT /oil-productions"),
    ("Produksi Minyak", "Detail, edit, hapus dan catatan", "Produksi tersedia", "Buka/edit/hapus; CRUD catatan; coba catatan invalid.", "Semua CRUD bekerja dan tervalidasi.", "/oil-productions/{id}"),
    ("Pembiayaan", "List, cari, filter, paginasi", "Transaksi tersedia", "Cari nama/deskripsi; filter petani/produk; ubah halaman.", "Hasil, total dan halaman konsisten.", "GET /financings"),
    ("Pembiayaan", "Create dan subtotal", "Petani/produk tersedia", "Isi transaksi valid; periksa harga produk/override; simpan.", "Subtotal = harga x quantity; relasi dan audit benar.", "POST /financings"),
    ("Pembiayaan", "Validasi, edit, dan hapus", "Transaksi tersedia", "Coba wajib kosong/tanggal invalid/harga negatif/qty <=0/UUID invalid; edit valid; hapus.", "Invalid ditolak; edit/hapus valid berhasil.", "PUT/DELETE /financings/{id}"),
    ("Penjualan", "List, cari, filter, paginasi", "Transaksi tersedia", "Cari/filter penjual, pembeli, produk, tanggal; ubah halaman.", "Hasil dan paginasi benar.", "GET /sales"),
    ("Penjualan", "Create, subtotal, edit dan hapus", "Produk/petani/mitra tersedia", "Buat transaksi valid; cek subtotal; edit; hapus.", "Subtotal = harga x quantity dan CRUD persisten.", "/sales"),
    ("Penjualan", "Validasi transaksi invalid", "Login", "Coba field wajib kosong, tanggal invalid, quantity <=0, harga negatif.", "Invalid ditolak tanpa record baru.", "POST/PUT /sales"),
    ("Demo", "Dashboard, laporan dan performa", "Mock data tersedia", "Buka /demo/dashboard, /petani, /produksi, /penjualan; uji filter/form.", "Data mock, filter, grafik, dan validasi form bekerja.", "Mock service"),
    ("Demo", "Operasional, inventory, kualitas", "Mock data tersedia", "Buka /operasional, /inventory, /kualitas; uji form dan notifikasi.", "Tampilan/form mock berjalan tanpa error.", "Mock service"),
    ("Ketahanan", "Loading, empty, error, retry", "Dapat simulasi respons", "Uji respons lambat, kosong, 4xx/5xx/network error dan Retry.", "State informatif, submit ganda dicegah, retry berfungsi.", "Semua API"),
    ("Keamanan", "API tanpa token", "Token dihapus", "Panggil endpoint terproteksi tanpa Authorization.", "Backend menolak 401/403; data tidak bocor/berubah.", "Semua API terproteksi"),
    ("Integritas", "Audit user_update", "Login sebagai user UAT", "Create/edit/delete data modul utama dan periksa audit.", "Identitas user login tercatat sesuai operasi.", "CRUD utama"),
    ("Kompatibilitas", "Responsive dan browser", "Chrome/Edge; desktop/tablet/mobile", "Uji navigasi, form, modal, chart, peta, tabel/card.", "Tidak ada konten penting terpotong dan seluruh aksi dapat dipakai.", "-"),
]

doc = Document()
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.5)
sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)
doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(10)
for name in ("Title", "Heading 1", "Heading 2"):
    doc.styles[name].font.name = "Arial"
    doc.styles[name].font.color.rgb = RGBColor(23, 107, 82)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(65)
r = p.add_run("DOKUMENTASI\nUSER ACCEPTANCE TEST (UAT)")
r.bold = True; r.font.name = "Arial"; r.font.size = Pt(24); r.font.color.rgb = RGBColor(23, 107, 82)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Sistem Informasi NILAM Manado")
r.bold = True; r.font.name = "Arial"; r.font.size = Pt(18)
p = doc.add_paragraph("Frontend: vue-nilam-manado\nBackend: fastapi-nilam-manado\nVersi 1.1 | 26 Agustus 2026")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph("\n")
table(doc, ["Disiapkan", "Ditinjau", "Disetujui"], [["Nama: __________", "Nama: __________", "Nama: __________"], ["Tanggal: ________", "Tanggal: ________", "Tanggal: ________"]], [5.5, 5.5, 5.5], 9)
doc.add_page_break()

heading(doc, "1. Informasi Dokumen")
table(doc, ["Atribut", "Nilai"], [["Nama", "Dokumentasi UAT NILAM Manado"], ["Versi", "1.1"], ["Tanggal", "26 Agustus 2026"], ["Sumber", "Kode frontend, backend, dokumentasi API, dan verifikasi implementasi QR"], ["Status", "Draft siap eksekusi"]], [4, 12.5], 9)
heading(doc, "2. Tujuan dan Ruang Lingkup")
doc.add_paragraph("Panduan dan bukti penerimaan pengguna untuk memverifikasi proses bisnis dari UI Vue hingga API/persistensi FastAPI, termasuk halaman mock/demo yang masih berada dalam aplikasi.")
bullet(doc, "Cakupan: autentikasi, dashboard, profil, wilayah, petani, lahan, mitra, produk, produksi dan catatan, pembiayaan, penjualan, laporan demo yang relevan, ketahanan UI, keamanan, audit, dan kompatibilitas.")
bullet(doc, "Di luar cakupan: QR dan traceability karena implementasi belum lengkap; penetration/load test mendalam; migrasi historis; dan integrasi eksternal yang tidak dipanggil frontend.")

heading(doc, "3. Fitur di Luar Ruang Lingkup UAT")
table(doc, ["Fitur", "Keputusan", "Alasan"], [
    ["Create/render QR", "Tidak diuji pada UAT rilis ini", "Implementasi masih berupa QR demo dengan traceCode dan quickLotId statis, belum dibentuk dari record backend."],
    ["Traceability lot", "Tidak diuji pada UAT rilis ini", "Daftar dan timeline masih berasal dari mockErpService, belum terintegrasi dengan backend FastAPI."],
    ["Scan QR dengan kamera", "Tidak diuji pada UAT rilis ini", "Komponen scanner, akses kamera, dan dependency pembaca QR/barcode belum tersedia."],
], [4, 5, 7.5], 8)

heading(doc, "4. Environment dan Data Uji")
table(doc, ["Item", "Isian"], [["URL Frontend", "____________"], ["Base URL Backend", "____________"], ["Commit Frontend/Backend", "____________"], ["Database", "____________"], ["Akun UAT", "____________"], ["Browser/Perangkat", "____________"], ["Tanggal Eksekusi", "____________"]], [5, 11.5], 9)
doc.add_paragraph("Data minimum: 2 akun, 3 petani lintas wilayah, 3 lahan, mitra, produk biaya/penjualan, produksi berstatus rencana/berjalan/selesai, dan transaksi lintas bulan. Gunakan prefix UAT-.")

heading(doc, "5. Kriteria Penerimaan")
bullet(doc, "Status kasus: Belum Diuji, Pass, Fail, atau Blocked. Fail wajib memiliki defect ID dan bukti.")
bullet(doc, "Exit: seluruh prioritas tinggi dieksekusi; tidak ada defect Critical/High terbuka; minimal 95% Pass; retest selesai; Product Owner sign-off.")
bullet(doc, "QR dan traceability baru dimasukkan ke UAT pada rilis berikutnya setelah create QR dinamis, sumber data backend, serta scan QR selesai diimplementasikan dan siap diuji end-to-end.")

heading(doc, "6. Kasus UAT")
doc.add_paragraph(f"Total {len(raw_cases)} kasus. Kolom eksekusi diisi oleh tester.")
for idx, (module, title, pre, steps, expected, api) in enumerate(raw_cases, 1):
    heading(doc, f"UAT-{idx:03d} – {module}: {title}", 2)
    table(doc, ["Atribut", "Detail"], [["Prasyarat", pre], ["Langkah", steps], ["Hasil Diharapkan", expected], ["Referensi", api], ["Hasil Aktual", ""], ["Status", "Belum Diuji"], ["Tester/Tanggal", ""], ["Bukti/Defect ID", ""]], [4, 12.5], 8)
    doc.add_paragraph("")

heading(doc, "7. Matriks Frontend–Backend")
table(doc, ["Fitur", "Route Frontend", "Backend/Sumber"], [
    ["Auth", "/real/login; route guard", "/auth/login, /auth/register"], ["Dashboard", "/real/dashboard", "/dashboard/*"],
    ["Petani & Wilayah", "/real/petani/*; /real/wilayah", "/farmers; /wilayah/*"], ["Lahan", "/real/lahan/*", "/lands; /lands/{id}/foto"],
    ["Produksi", "/real/produksi-tanam/*; /real/produksi-minyak/*", "/planting-productions; /oil-productions; notes"],
    ["Mitra & Produk", "/real/mitra; /real/produk-*", "/partners; /financing-products; /sales-products"],
    ["Transaksi", "/real/pembiayaan; /real/penjualan", "/financings; /sales"],
], [4, 6, 6.5], 8)

heading(doc, "8. Log Defect dan Rekap")
table(doc, ["Defect ID", "UAT ID", "Ringkasan", "Severity", "Status", "Owner", "Bukti"], [["", "", "", "", "", "", ""] for _ in range(8)], [2, 1.8, 4, 2, 2, 2, 3], 7)
doc.add_paragraph("")
table(doc, ["Metrik", "Jumlah"], [["Total", len(raw_cases)], ["Pass", ""], ["Fail", ""], ["Blocked", ""], ["Belum Diuji", len(raw_cases)], ["Pass Rate", "____ %"]], [8, 8.5], 9)

heading(doc, "9. Persetujuan UAT")
doc.add_paragraph("Keputusan: ☐ Diterima  ☐ Diterima dengan catatan  ☐ Ditolak")
table(doc, ["Peran", "Nama", "Tanda Tangan", "Tanggal", "Catatan"], [["Key User", "", "", "", ""], ["UAT Lead", "", "", "", ""], ["Product Owner", "", "", "", ""], ["Developer/QA", "", "", "", ""]], [3, 3, 3.5, 2.5, 4.5], 8)

for section in doc.sections:
    hp = section.header.paragraphs[0]
    hp.text = "NILAM Manado | Dokumentasi UAT | Versi 1.1"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Halaman ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)

doc.core_properties.title = "Dokumentasi UAT NILAM Manado"
doc.core_properties.subject = "UAT frontend Vue dan backend FastAPI termasuk status QR"
doc.core_properties.author = "Tim NILAM Manado"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"Created: {OUT} ({len(raw_cases)} cases)")
